#!/usr/bin/env python3
"""生成 P-ADM-04 批量导入所需的真实 .xlsx 文件。

产出（相对仓库根目录）：
  assets/people-import-template.xlsx   下载用导入模板（数据填写页 + 填写说明页）
  assets/samples/*.xlsx                示例文件，覆盖 M01 §4.2 的文件级与行级判定

依赖：openpyxl（仅生成时需要，原型运行时不依赖任何库）
用法：python3 tools/gen-import-template.py
"""

import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(ROOT, "assets")
SAMPLES = os.path.join(ASSETS, "samples")

DATA_SHEET = "数据填写页"
HELP_SHEET = "填写说明页"
HEADERS = ["姓名", "邮箱", "密码", "手机号", "角色"]
ROLES = ["销售", "销售主管", "解决方案人员", "超级管理员"]
MAX_ROWS = 500          # 单批行数上限（建议初值，待联调核定）
DV_ROWS = MAX_ROWS + 1  # 数据校验覆盖到第 501 行

BRAND = "2C5CE6"
GREY = "6B7280"
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="D8DEE9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COL_RULES = {
    "姓名": (
        "必填",
        "人员真实姓名。",
        "留空该行判定为不合法，不予导入。",
    ),
    "邮箱": (
        "必填",
        "需填写完整邮箱地址，后缀固定为 @teddymobile.cn。\n"
        "公司邮箱命名规则为姓名全拼小写，例如「张三」为 zhangsan@teddymobile.cn。\n"
        "系统以本人实际企业邮箱为准，不校验前缀与姓名是否一致。",
        "只填前缀不填后缀、缺少 @、使用其他公司或个人邮箱后缀、"
        "前缀含空格或中文，均判定为不合法。邮箱在系统内唯一，重复不予导入。",
    ),
    "密码": (
        "必填",
        "明文填写人员初始登录密码，系统不代为生成。",
        "留空该行判定为不合法。人员首次登录不要求强制修改密码。",
    ),
    "手机号": (
        "必填",
        "11 位手机号。",
        "格式不正确判定为不合法。",
    ),
    "角色": (
        "可选",
        "在单元格下拉中直接选择，仅限：" + "、".join(ROLES) + "。",
        "手工输入且超出上述范围判定为不合法。留空可以导入，"
        "该人员进入待分配角色池、账号为停用状态，分配角色后仍需超级管理员手动启用。",
    ),
}

BATCH_RULES = [
    "单次导入上限为 %d 行、5MB（建议初值，最终数值待联调确认）。超限时请拆批上传，系统不会静默截断。" % MAX_ROWS,
    "上传后系统先做文件级检测：非模板格式、模板列不完整、数据填写页无数据行，均直接提示且不进入逐行检测。",
    "逐行检测时，一行存在多个问题会一次性全部列出，不需要反复修正、反复上传。",
    "不合法行不会阻断同一批次中合法行的导入；不合法数据可导出、修正后重新上传。",
    "同一文件内邮箱重复的，重复行全部判定为不合法，需人工确认保留哪一条。",
    "邮箱在去除首尾空格并统一按小写处理后再校验与判重，不因大小写或空格差异被识别为两个账号。",
    "确认导入后若过程中断，已创建的部分全部撤销，本批次视为未导入，需重新上传。",
    "本导入仅用于创建新账号，不支持通过导入更新已有人员信息；模板不含钉钉身份列，需在人员详情逐个维护。",
]

EXAMPLE = ["张三", "zhangsan@teddymobile.cn", "Tk8m#Wp3Qz6v", "13800000000", "销售"]


def style_header(ws, headers=None, note=True):
    """表头：必填列品牌色、可选列灰色，附单元格批注。"""
    headers = headers or HEADERS
    widths = {"姓名": 14, "邮箱": 34, "密码": 20, "手机号": 16, "角色": 18}
    for i, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = HEAD_FONT
        required = COL_RULES.get(name, ("可选",))[0] == "必填"
        cell.fill = PatternFill("solid", fgColor=BRAND if required else GREY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 18)
        if note and name in COL_RULES:
            flag, rule, err = COL_RULES[name]
            body = "%s（%s）\n%s\n%s" % (name, flag, rule, err)
            cell.comment = Comment(body, "销售指挥官", height=190, width=340)
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"


def add_validations(ws):
    """角色列内置固定枚举下拉；其余列挂选中提示。"""
    role_col = get_column_letter(HEADERS.index("角色") + 1)
    dv_role = DataValidation(
        type="list",
        formula1='"%s"' % ",".join(ROLES),
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="角色超出范围",
        error="角色只能是：" + "、".join(ROLES) + "。留空则进入待分配角色池。",
        promptTitle="角色（可选）",
        prompt="点击右侧箭头直接选择。留空可以导入，该人员进入待分配角色池、账号停用。",
        showInputMessage=True,
    )
    ws.add_data_validation(dv_role)
    dv_role.add("%s2:%s%d" % (role_col, role_col, DV_ROWS))

    prompts = {
        "姓名": ("姓名（必填）", "人员真实姓名，留空该行不合法。"),
        "邮箱": (
            "邮箱（必填）",
            "填写完整邮箱地址，后缀固定 @teddymobile.cn；命名规则为姓名全拼小写"
            "（张三 → zhangsan@teddymobile.cn），以本人实际企业邮箱为准。",
        ),
        "密码": ("密码（必填）", "明文填写初始密码，系统不代为生成，留空该行不合法。"),
        "手机号": ("手机号（必填）", "11 位手机号，格式不正确该行不合法。"),
    }
    for name, (title, text) in prompts.items():
        col = get_column_letter(HEADERS.index(name) + 1)
        dv = DataValidation(
            type="textLength",
            operator="lessThanOrEqual",
            formula1=254,
            allow_blank=True,
            showInputMessage=True,
            promptTitle=title,
            prompt=text,
        )
        ws.add_data_validation(dv)
        dv.add("%s2:%s%d" % (col, col, DV_ROWS))


def build_help(ws):
    """填写说明页：各列规则 + 批次规则 + 一行正确示例。说明页不参与检测。"""
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 62

    title = Font(name="微软雅黑", size=13, bold=True, color="1F2937")
    section = Font(name="微软雅黑", size=11, bold=True, color=BRAND)
    wrap = Alignment(vertical="top", wrap_text=True)

    ws["A1"] = "人员批量导入模板 · 填写说明"
    ws["A1"].font = title
    ws["A2"] = "本页仅为说明，不参与检测；请在【%s】填写数据。数据填写页只有表头、不预置示例行。" % DATA_SHEET
    ws["A2"].font = Font(name="微软雅黑", size=10, color="B45309")

    r = 4
    ws.cell(row=r, column=1, value="一、各列填写规则").font = section
    r += 1
    for i, name in enumerate(["列名", "是否必填", "填写规则", "常见错误与后果"], start=1):
        c = ws.cell(row=r, column=i, value=name)
        c.font = HEAD_FONT
        c.fill = PatternFill("solid", fgColor=BRAND)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    r += 1
    for name in HEADERS:
        flag, rule, err = COL_RULES[name]
        for i, v in enumerate([name, flag, rule, err], start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = BODY_FONT
            c.alignment = wrap
            c.border = BOX
        ws.row_dimensions[r].height = 46
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="二、批次与检测规则").font = section
    r += 1
    for line in BATCH_RULES:
        c = ws.cell(row=r, column=1, value="· " + line)
        c.font = BODY_FONT
        c.alignment = wrap
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 20
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="三、正确格式示例（示例仅在本页，不要粘贴到数据填写页）").font = section
    r += 1
    for i, name in enumerate(HEADERS, start=1):
        c = ws.cell(row=r, column=i, value=name)
        c.font = HEAD_FONT
        c.fill = PatternFill("solid", fgColor=GREY)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    r += 1
    for i, v in enumerate(EXAMPLE, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = BODY_FONT
        c.border = BOX


def new_book(headers=None, sheet=DATA_SHEET, note=True, validations=True):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    style_header(ws, headers, note=note)
    if validations and (headers is None or headers == HEADERS):
        add_validations(ws)
    return wb, ws


def write_rows(ws, rows):
    for n, row in enumerate(rows, start=2):
        for i, v in enumerate(row, start=1):
            if v in (None, ""):
                continue
            c = ws.cell(row=n, column=i, value=v)
            c.font = BODY_FONT


def save(wb, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print("  %-46s %6.1f KB" % (os.path.relpath(path, ROOT), os.path.getsize(path) / 1024))


def gen_template():
    wb, _ = new_book()
    build_help(wb.create_sheet(HELP_SHEET))
    save(wb, os.path.join(ASSETS, "people-import-template.xlsx"))


def gen_samples():
    """示例文件全部为真实 .xlsx，走原型内同一套校验。"""

    # 混合结果：3 行合法、3 行不合法
    wb, ws = new_book()
    build_help(wb.create_sheet(HELP_SHEET))
    write_rows(ws, [
        ["何书桐", "heshutong@teddymobile.cn", "Hs7k#Qm2Xw9p", "13902212205", None],
        ["徐澈", "xuche@teddymobile.cn", "Xc4d#Kp8Ze1m", "13803318899", "销售"],
        ["罗一鸣", "luoyiming@teddymobile.cn", "Ly6w#Nt4Bh7c", "13577412266", "解决方案人员"],
        ["谢知微", "zhiwei@gmail.com", "Zw3m#Fj9Vd5r", "15899643300", "销售"],
        [None, "fuzan@teddymobile.cn", "Fz8n#Rt5Cq2k", "13555106677", "销售王"],
        ["蒋斯年", "jiangsinian@teddymobile.cn", None, "手机号", "销售"],
    ])
    save(wb, os.path.join(SAMPLES, "sample-mixed.xlsx"))

    # 同一文件内邮箱重复：重复行均不合法，并演示去空格 + 统一小写后判重
    wb, ws = new_book()
    write_rows(ws, [
        ["汪清明", "wangqingming@teddymobile.cn", "Wq5t#Hn3Bd8v", "13902212288", "销售"],
        ["于书白", "yushubai@teddymobile.cn", "Ys2k#Lm7Pc4x", "13588170099", "销售"],
        ["缪清", "miaoqing@teddymobile.cn", "Mq9r#Vd6Kt3n", "13577412299", "解决方案人员"],
        ["于书白（新）", "  YuShuBai@teddymobile.CN  ", "Ys4w#Nt8Bh2c", "13844601122", "销售主管"],
    ])
    save(wb, os.path.join(SAMPLES, "sample-dup.xlsx"))

    # 全部行不合法：不产生任何账号
    wb, ws = new_book()
    write_rows(ws, [
        [None, None, None, None, "销售"],
        ["谢知微", "zhiwei@gmail.com", "Zw3m#Fj9Vd5r", "15899643300", "销售"],
        ["付赞", "fuzan@teddymobile.cn", "Fz8n#Rt5Cq2k", "13555106677", "销售"],
        ["蒋斯年", "jiangsinian", None, "手机号", "销售总监"],
    ])
    save(wb, os.path.join(SAMPLES, "sample-allbad.xlsx"))

    # 文件级：非模板格式（找不到数据填写页）
    wb, ws = new_book(headers=["员工姓名", "所属部门", "入职日期", "工号"],
                      sheet="Sheet1", note=False, validations=False)
    write_rows(ws, [["何书桐", "销售部", "2026-09-01", "E20260901"]])
    save(wb, os.path.join(SAMPLES, "sample-not-template.xlsx"))

    # 文件级：模板列不完整（缺密码、手机号）
    wb, ws = new_book(headers=["姓名", "邮箱", "角色"], validations=False)
    write_rows(ws, [["何书桐", "heshutong@teddymobile.cn", "销售"]])
    save(wb, os.path.join(SAMPLES, "sample-missing-cols.xlsx"))

    # 文件级：数据填写页无数据行（保留填写说明页，说明页不参与检测）
    wb, _ = new_book()
    build_help(wb.create_sheet(HELP_SHEET))
    save(wb, os.path.join(SAMPLES, "sample-empty.xlsx"))

    # 文件级：超出单批行数上限，提示拆批、不截断
    wb, ws = new_book()
    write_rows(ws, [
        ["演示%03d" % i, "demo%03d@teddymobile.cn" % i, "Dm%04d#Xk2p" % i,
         "138%08d" % i, "销售"]
        for i in range(1, MAX_ROWS + 21)
    ])
    save(wb, os.path.join(SAMPLES, "sample-oversize.xlsx"))


if __name__ == "__main__":
    print("生成导入模板：")
    gen_template()
    print("生成示例文件：")
    gen_samples()
